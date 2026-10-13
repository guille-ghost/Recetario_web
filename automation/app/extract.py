import io
from pathlib import Path

import pypdfium2 as pdfium
import pytesseract
from docx import Document
from pypdf import PdfReader

from .settings import Settings

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TEXT_MIMES = {"text/plain", "text/markdown", "text/csv", "text/tab-separated-values"}


def _decode_text(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        return content.decode("cp1252", errors="replace")


def extract_text(name: str, mime_type: str, content: bytes, settings: Settings) -> str:
    suffix = Path(name).suffix.lower()
    if mime_type in TEXT_MIMES or suffix in {".txt", ".md", ".csv", ".tsv"}:
        text = _decode_text(content)
    elif mime_type == DOCX_MIME or suffix == ".docx":
        document = Document(io.BytesIO(content))
        parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)
    elif mime_type == "application/pdf" or suffix == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        document = pdfium.PdfDocument(content)
        if len(reader.pages) > 100:
            document.close()
            raise ValueError("PDF exceeds the 100-page processing limit")
        pages = []
        for index, page in enumerate(reader.pages):
            page_text = (page.extract_text() or "").strip()
            if len(page_text) < 40:
                try:
                    bitmap = document[index].render(scale=200 / 72)
                    image = bitmap.to_pil()
                    try:
                        recognized = pytesseract.image_to_string(image, lang="spa").strip()
                    finally:
                        image.close()
                        bitmap.close()
                    if recognized:
                        page_text = recognized
                except Exception as error:
                    document.close()
                    raise RuntimeError("No se pudo ejecutar OCR en el PDF. Comprueba que Tesseract y el idioma español estén instalados.") from error
            if page_text:
                pages.append(page_text)
        document.close()
        text = "\n".join(pages)
        if not text.strip():
            raise ValueError("No se reconoció texto en el PDF. Comprueba la calidad de las imágenes.")
    else:
        raise ValueError(f"Unsupported file type: {mime_type or suffix or 'unknown'}")

    normalized = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not normalized:
        raise ValueError("No readable text found in the file")
    if len(normalized) > settings.max_text_chars:
        normalized = normalized[: settings.max_text_chars]
    return normalized
